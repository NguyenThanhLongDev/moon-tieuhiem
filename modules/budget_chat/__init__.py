"""Budget Chat — báo ngân sách FB Ads thay Zalo.

Routes:
  GET  /ngan-sach                       → redirect đến team của user
  GET  /ngan-sach/team/<team_id>        → chat room
  GET  /ngan-sach/messages              → poll JSON {messages, items}
  POST /ngan-sach/send                  → AI parse + save
  GET  /ngan-sach/tong-hop              → manager dashboard
  POST /ngan-sach/items/<id>/edit       → sửa item
  POST /ngan-sach/items/<id>/delete     → xoá item
  GET  /ngan-sach/api/match-tk?q=       → search TK QC để gán vào item

Permission:
  staff/sale/leader → vào team mình (R+W)
  manager/admin/superadmin → mọi team + tab Tổng hợp
  accountant → read tab Tổng hợp
"""
from __future__ import annotations

import logging
from datetime import date, datetime, timedelta
from functools import wraps
from typing import Optional

from flask import (
    Blueprint, jsonify, redirect, render_template, request, session, url_for, flash, abort,
)

logger = logging.getLogger(__name__)

budget_chat_bp = Blueprint(
    "budget_chat", __name__,
    url_prefix="/ngan-sach",
    template_folder="templates",
)

# Roles
_NV_ROLES = {"sale", "leader", "staff", "kho", "manager", "admin", "superadmin", "it", "accountant", "ketoan"}
_MANAGER_ROLES = {"manager", "admin", "superadmin"}
_READONLY_SUMMARY_ROLES = {"accountant", "ketoan"}


def _current_user() -> dict:
    """Trả {id, username, full_name, role, team_id (int), team_code (str)} của user."""
    uid = session.get("user_id")
    if not uid:
        return {}
    try:
        from db import get_conn
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute("""
                    SELECT u.id, u.username, u.full_name, u.role,
                           u.team_id, t.team_code, t.team_name
                      FROM users u
                      LEFT JOIN teams t ON t.id = u.team_id
                     WHERE u.id = %s
                """, (int(uid),))
                r = cur.fetchone()
                if not r:
                    return {}
                return {
                    "id": int(r[0]),
                    "username": r[1] or "",
                    "full_name": r[2] or r[1] or "",
                    "role": (r[3] or "").lower(),
                    "team_id": int(r[4]) if r[4] else None,
                    "team_code": r[5] or "",
                    "team_name": r[6] or "",
                }
    except Exception as e:
        logger.warning("_current_user error: %s", e)
        return {}


def _login_required(f):
    @wraps(f)
    def wrapper(*a, **kw):
        if not session.get("user_id"):
            return redirect(url_for("auth.login"))
        return f(*a, **kw)
    return wrapper


def _can_view_team(user: dict, team_code: str) -> bool:
    """Manager + kế toán xem tất; NV thường chỉ xem team mình.

    Kế toán vào chat để kiểm tra Lan parse đúng/sai → giữ quyền view all.
    """
    role = (user.get("role") or "").lower()
    if role in _MANAGER_ROLES or role in _READONLY_SUMMARY_ROLES:
        return True
    user_team = (user.get("team_code") or "")
    return bool(user_team) and user_team == team_code


def _can_write_team(user: dict, team_code: str) -> bool:
    """Manager + NV team đó mới được gửi tin. Kế toán read-only (xem ko chat)."""
    role = (user.get("role") or "").lower()
    if role in _MANAGER_ROLES:
        return True
    if role in _READONLY_SUMMARY_ROLES:
        return False
    user_team = (user.get("team_code") or "")
    return role in _NV_ROLES and bool(user_team) and user_team == team_code


def _can_view_summary(user: dict) -> bool:
    role = (user.get("role") or "").lower()
    return role in (_MANAGER_ROLES | _READONLY_SUMMARY_ROLES)


def _push_lan_to_zalo(team_code: str, text: str, user_id: Optional[int] = None) -> None:
    """Đẩy tin Lan từ DB ra Zalo group qua bridge HTTP local.

    Lookup thread_id từ app_config['zalo_thread_<tid>'] = team_code (reverse).
    Nếu có user_id, tag @<full_name> ở đầu tin với mention chuẩn Zalo.
    No-op nếu team_code chưa được map. Fire-and-forget, không raise.
    """
    if not text or not team_code:
        return
    try:
        import os as _os, requests as _rq
        from app_ctx import load_config
        secret = (_os.environ.get("ZALO_BRIDGE_SECRET") or "").strip()
        if not secret:
            return
        cfg = load_config()
        thread_id = None
        for k, v in cfg.items():
            if isinstance(k, str) and k.startswith("zalo_thread_") and str(v).strip() == team_code:
                thread_id = k[len("zalo_thread_"):]
                break
        if not thread_id:
            return

        # Build mention nếu có user_id + user có zalo_uid
        payload = {"thread_id": thread_id, "text": text}
        if user_id:
            try:
                from db import get_conn
                with get_conn() as conn:
                    with conn.cursor() as cur:
                        cur.execute(
                            "SELECT zalo_uid, full_name FROM users WHERE id=%s AND zalo_uid IS NOT NULL",
                            (int(user_id),),
                        )
                        row = cur.fetchone()
                if row and row[0] and row[1]:
                    zalo_uid, full_name = row[0], row[1]
                    tag = f"@{full_name} "
                    payload["text"] = tag + text
                    # pos=0, len=UTF-16 length của "@<name>" — Zalo dùng UTF-16
                    tag_no_space = f"@{full_name}"
                    utf16_len_tag = sum(2 if ord(c) > 0xFFFF else 1 for c in tag_no_space)
                    payload["mentions"] = [{"pos": 0, "uid": zalo_uid, "len": utf16_len_tag}]
            except Exception as exc:
                logger.info("push_lan_to_zalo mention build skip (%s)", exc)

        url = _os.environ.get("ZALO_BRIDGE_OUTBOUND_URL", "http://127.0.0.1:5051/send")
        _rq.post(url, json=payload, headers={"X-Bridge-Secret": secret}, timeout=3)
    except Exception as exc:
        logger.info("push_lan_to_zalo skip (%s)", exc)


def _is_telegram_enabled() -> bool:
    """Toggle forward chat NS sang Telegram. Default BẬT."""
    try:
        from app_ctx import load_config
        v = load_config().get("budget_chat_telegram_enabled")
        if v is None:
            return True
        return str(v).strip().lower() in ("1", "true", "yes", "on", "bật", "bat")
    except Exception:
        return True


def _list_all_teams() -> list[dict]:
    """Load tất cả team active từ bảng teams (join với users đếm thành viên)."""
    try:
        from db import get_conn
        with get_conn() as conn:
            with conn.cursor() as cur:
                # ⚠ Chỉ team_type = 'kinh_doanh' mới phải báo ngân sách FB Ads.
                # Team sale (Trọng Nam) + team khác (kho) → loại trừ.
                cur.execute("""
                    SELECT t.team_code, t.team_name, COUNT(u.id) AS n
                      FROM teams t
                      LEFT JOIN users u ON u.team_id = t.id
                     WHERE t.status = 'active'
                       AND t.team_type = 'kinh_doanh'
                     GROUP BY t.team_code, t.team_name
                     ORDER BY t.team_name
                """)
                return [{
                    "team_id": r[0],     # = team_code, dùng làm URL param
                    "name": r[1] or (r[0] or "").replace("team-", "Team ").title(),
                    "member_count": int(r[2] or 0),
                } for r in cur.fetchall()]
    except Exception as e:
        logger.warning("_list_all_teams error: %s", e)
        return []


def _team_code_to_int_id(team_code: str) -> Optional[int]:
    """team_code (vd 'team-nam') → teams.id (int). Cache trong app context tuỳ chọn."""
    if not team_code:
        return None
    try:
        from db import get_conn
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT id FROM teams WHERE team_code=%s", (team_code,))
                r = cur.fetchone()
                return int(r[0]) if r else None
    except Exception:
        return None


@budget_chat_bp.route("/")
@_login_required
def index():
    user = _current_user()
    if not user:
        return redirect(url_for("auth.login"))
    role = (user.get("role") or "").lower()
    # Manager → tổng hợp luôn
    if role in _MANAGER_ROLES:
        return redirect(url_for("budget_chat.summary"))
    # NV thường → team mình
    team_code = user.get("team_code") or ""
    if not team_code:
        flash("Bạn chưa được gán team. Liên hệ admin.", "warning")
        return redirect("/")
    return redirect(url_for("budget_chat.team_room", team_id=team_code))


@budget_chat_bp.route("/team/<team_id>")
@_login_required
def team_room(team_id: str):
    user = _current_user()
    if not _can_view_team(user, team_id):
        abort(403)
    # Chặn team không phải kinh_doanh — báo NS chỉ áp cho team kinh doanh.
    try:
        from db import get_conn
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT team_type FROM teams WHERE team_code=%s", (team_id,))
                row = cur.fetchone()
                if not row:
                    flash(f"Team '{team_id}' không tồn tại.", "danger")
                    return redirect(url_for("budget_chat.index"))
                if row[0] != "kinh_doanh":
                    flash(f"Team này không phải team kinh doanh — không áp dụng báo NS.", "warning")
                    return redirect(url_for("budget_chat.index"))
    except Exception as e:
        logger.warning("team_type check error: %s", e)
    teams = _list_all_teams() if (user.get("role") or "").lower() in _MANAGER_ROLES else []
    can_write = _can_write_team(user, team_id)
    return render_template(
        "budget_chat/room.html",
        current_user=user,
        team_id=team_id,
        team_name=team_id.replace("team-", "Team ").title(),
        teams=teams,
        can_write=can_write,
        is_manager=(user.get("role") or "").lower() in _MANAGER_ROLES,
    )


@budget_chat_bp.route("/messages")
@_login_required
def messages_poll():
    user = _current_user()
    team_id = request.args.get("team_id", "").strip()
    after_id = int(request.args.get("after_id", "0") or 0)
    if not _can_view_team(user, team_id):
        return jsonify({"error": "forbidden"}), 403

    from db import get_conn
    from modules.budget_chat.repository import list_messages, list_items_for_message, mark_read

    with get_conn() as conn:
        with conn.cursor() as cur:
            msgs = list_messages(cur, team_id, limit=200, after_id=after_id)
            # Đính kèm items per message
            for m in msgs:
                m["items"] = list_items_for_message(cur, m["id"])
            # Mark read up to newest msg ID
            if msgs:
                top_id = max(m["id"] for m in msgs)
                mark_read(cur, int(user["id"]), team_id, top_id)
                conn.commit()

    return jsonify({"messages": msgs})


@budget_chat_bp.route("/send", methods=["POST"])
@_login_required
def send_message():
    user = _current_user()
    team_id = request.form.get("team_id", "").strip()
    body = (request.form.get("body") or "").strip()
    if not body:
        return jsonify({"error": "empty body"}), 400
    if not _can_write_team(user, team_id):
        return jsonify({"error": "forbidden"}), 403

    from db import get_conn
    from modules.budget_chat.repository import (
        insert_message, save_parse_result, insert_lan_message,
        fill_missing_date_for_user,
    )
    from modules.budget_chat.ai_parser import parse_budget_message
    from modules.budget_chat.lan_personality import (
        is_budget_intent, ai_reply_parse_fail, ai_reply_ask_date,
        ai_reply_fill_date, extract_fill_date_intent,
    )
    try:
        from tz_utils import now_hcm
        _today = now_hcm().date()
    except Exception:
        from datetime import date as _date
        _today = _date.today()
    from datetime import timedelta as _td
    today_iso = _today.isoformat()
    tomorrow_iso = (_today + _td(days=1)).isoformat()

    # 1. Insert message ngay (NV thấy hiển thị tức thì)
    with get_conn() as conn:
        with conn.cursor() as cur:
            msg_id = insert_message(cur, team_id, int(user["id"]), body)
            conn.commit()

    # 1B. Intent "fill ngày": NV chat tin ngắn chỉ chứa ngày → fill items
    # thiếu ngày của chính NV này trong team trong 24h gần đây.
    fill_date = extract_fill_date_intent(body, today_iso, tomorrow_iso)
    if fill_date:
        nv_name = user.get("full_name") or user.get("username") or "bạn"
        with get_conn() as conn:
            with conn.cursor() as cur:
                filled = fill_missing_date_for_user(
                    cur, int(user["id"]), team_id, fill_date, int(user["id"])
                )
                lan_reply_id = None
                if filled:
                    from datetime import date as _d
                    try:
                        for_date_vn = _d.fromisoformat(fill_date).strftime("%d/%m/%Y")
                    except Exception:
                        for_date_vn = fill_date
                    summary = ", ".join(
                        f"TK {it['tk_name_raw']}"
                        + (f" thẻ {it['card_last4']}" if it['card_last4'] else "")
                        + f" {it['amount_vnd']:,}đ".replace(",", ".")
                        for it in filled[:5]
                    )
                    lan_text = ai_reply_fill_date(nv_name, for_date_vn, summary)
                    lan_reply_id = insert_lan_message(cur, team_id, lan_text)
                conn.commit()
        if filled:
            _push_lan_to_zalo(team_id, lan_text, user_id=int(user["id"]))
            return jsonify({
                "ok": True, "message_id": msg_id,
                "filled_count": len(filled), "lan_reply_id": lan_reply_id,
            })
        # Không có items thiếu ngày → fall-through xuống parse như tin chat thường

    # 2. Parse AI (sync — DeepSeek thường <2s)
    parsed = parse_budget_message(body)

    # 3. Save items + nếu parse fail nhưng có ý báo NS → Lan reply tag tên NV
    items_count = len(parsed.get("items") or [])
    lan_reply_id = None
    lan_text_for_zalo = None
    with get_conn() as conn:
        with conn.cursor() as cur:
            save_parse_result(cur, msg_id, parsed)
            nv_name = user.get("full_name") or user.get("username") or "bạn"
            if items_count == 0 and is_budget_intent(body):
                lan_text = ai_reply_parse_fail(nv_name, body)
                lan_reply_id = insert_lan_message(cur, team_id, lan_text)
                lan_text_for_zalo = lan_text
            elif items_count > 0 and not parsed.get("for_date"):
                # Có TK + tiền nhưng thiếu ngày → đã lưu items với for_date NULL,
                # Lan hỏi lại để NV bổ sung qua nút "Sửa ngày".
                items_summary = ", ".join(
                    f"TK {it.get('tk_name','?')}"
                    + (f" thẻ {it.get('card_last4')}" if it.get('card_last4') else "")
                    + f" {int(it.get('amount_vnd') or 0):,}đ".replace(",", ".")
                    for it in (parsed.get("items") or [])[:5]
                )
                lan_text = ai_reply_ask_date(nv_name, body, items_summary)
                lan_reply_id = insert_lan_message(cur, team_id, lan_text)
                lan_text_for_zalo = lan_text
            elif items_count > 0 and parsed.get("for_date"):
                # Happy path: NS đầy đủ items + ngày → xác nhận ghi nhận
                items_summary = ", ".join(
                    f"TK {it.get('tk_name','?')}"
                    + (f" thẻ {it.get('card_last4')}" if it.get('card_last4') else "")
                    + f" {int(it.get('amount_vnd') or 0):,}đ".replace(",", ".")
                    for it in (parsed.get("items") or [])[:5]
                )
                try:
                    from datetime import date as _d
                    _fd = _d.fromisoformat(parsed["for_date"])
                    for_date_vn = _fd.strftime("%d/%m/%Y")
                except Exception:
                    for_date_vn = str(parsed.get("for_date") or "")
                from modules.budget_chat.lan_personality import ai_reply_ack_edit
                lan_text = ai_reply_ack_edit(nv_name, items_summary, for_date_vn)
                lan_reply_id = insert_lan_message(cur, team_id, lan_text)
                lan_text_for_zalo = lan_text
            conn.commit()
    if lan_text_for_zalo:
        _push_lan_to_zalo(team_id, lan_text_for_zalo, user_id=int(user["id"]))

    # 4. Telegram bridge (non-blocking — log fail không raise)
    if _is_telegram_enabled():
        try:
            from modules.budget_chat.telegram_bridge import notify_new_message
            notify_new_message(user, team_id, body, parsed)
        except Exception as e:
            logger.warning("notify_new_message error: %s", e)
    else:
        logger.info("Telegram forward TẮT — skip notify_new_message msg_id=%s", msg_id)

    return jsonify({
        "ok": True,
        "message_id": msg_id,
        "items_count": items_count,
        "lan_reply_id": lan_reply_id,
    })


@budget_chat_bp.route("/tong-hop")
@_login_required
def summary():
    user = _current_user()
    if not _can_view_summary(user):
        abort(403)
    # Default: ngày mai
    try:
        from tz_utils import now_hcm
        today = now_hcm().date()
    except Exception:
        today = date.today()
    default_date = request.args.get("date") or (today + timedelta(days=1)).isoformat()
    try:
        target_date = date.fromisoformat(default_date)
    except Exception:
        target_date = today + timedelta(days=1)
    team_filter = request.args.get("team") or ""

    from db import get_conn
    from modules.budget_chat.repository import (
        summary_by_team, summary_by_user, list_users_should_report,
        summary_tk_type, list_items_by_tk_type,
    )
    with get_conn() as conn:
        with conn.cursor() as cur:
            teams_summary = summary_by_team(cur, target_date)
            users_summary = summary_by_user(cur, target_date, team_filter or None)
            should_report = list_users_should_report(cur, target_date, team_filter or None)
            tk_type_summary = summary_tk_type(cur, target_date)
            items_cty = list_items_by_tk_type(cur, target_date, "cty", team_filter or None)
            items_hkd = list_items_by_tk_type(cur, target_date, "hkd", team_filter or None)

    # Build maps
    reported_uids = {u["user_id"] for u in users_summary}
    not_reported = [u for u in should_report if not u["reported"]]

    total_all = sum(t["total"] for t in teams_summary)
    teams_all = _list_all_teams()
    teams_map = {t["team_id"]: t for t in teams_all}

    return render_template(
        "budget_chat/summary.html",
        current_user=user,
        target_date=target_date.isoformat(),
        team_filter=team_filter,
        teams_summary=teams_summary,
        users_summary=users_summary,
        not_reported=not_reported,
        total_all=total_all,
        teams_all=teams_all,
        teams_map=teams_map,
        should_count=len(should_report),
        reported_count=len(reported_uids),
        tk_type_summary=tk_type_summary,
        items_cty=items_cty,
        items_hkd=items_hkd,
        telegram_enabled=_is_telegram_enabled(),
    )


@budget_chat_bp.route("/export-excel")
@_login_required
def export_excel():
    """Xuất Excel ngân sách theo khoảng ngày.

    Params: ?date_from=YYYY-MM-DD&date_to=YYYY-MM-DD&team=<team_code>
    team rỗng → tất cả team (mỗi team 1 sheet).
    """
    user = _current_user()
    if not _can_view_summary(user):
        abort(403)

    try:
        from tz_utils import now_hcm
        today = now_hcm().date()
    except Exception:
        today = date.today()

    df = request.args.get("date_from") or today.isoformat()
    dt = request.args.get("date_to") or today.isoformat()
    team_filter = (request.args.get("team") or "").strip()
    try:
        date_from = date.fromisoformat(df)
        date_to = date.fromisoformat(dt)
    except Exception:
        date_from = date_to = today
    if date_to < date_from:
        date_from, date_to = date_to, date_from

    from db import get_conn
    from modules.budget_chat.repository import list_items_for_export
    with get_conn() as conn:
        with conn.cursor() as cur:
            items = list_items_for_export(cur, date_from, date_to, team_filter or None)

    # Build Excel
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    import io
    from datetime import timedelta as _td

    wb = Workbook()
    wb.remove(wb.active)

    # Build day list: từ date_from đến date_to inclusive
    days: list[date] = []
    d = date_from
    while d <= date_to:
        days.append(d)
        d += _td(days=1)

    # Group items theo team (1 sheet/team nếu xuất tất cả; 1 sheet duy nhất nếu chọn team)
    teams_map: dict[str, list[dict]] = {}
    team_names: dict[str, str] = {}
    for it in items:
        tid = it["team_id"]
        teams_map.setdefault(tid, []).append(it)
        team_names[tid] = it["team_name"] or tid
    if not teams_map:
        teams_map["_"] = []
        team_names["_"] = "Trống"

    # Style
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill = PatternFill("solid", fgColor="4F81BD")
    info_fill = PatternFill("solid", fgColor="DCE6F1")
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    money_fill = PatternFill("solid", fgColor="FDE9D9")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center")
    thin = Side(border_style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for tid, team_items in teams_map.items():
        team_label = team_names.get(tid, tid).replace("team-", "Team ").title() or "Team"
        ws = wb.create_sheet(title=(team_label[:28] or "Team"))

        # Aggregate theo (tk_name_raw, card_last4) → tổng + per-day
        # Lưu set NV + team để hiển thị (1 TK có thể nhiều NV báo trong range)
        agg: dict[tuple, dict] = {}
        for it in team_items:
            key = (it["tk_name_raw"], it["card_last4"])
            entry = agg.setdefault(key, {
                "tk": it["tk_name_raw"],
                "card": it["card_last4"],
                "who_set": [],     # list giữ thứ tự, dedup tay
                "team_set": [],
                "total": 0,
                "by_day": {},      # date_iso → amount
                "missing": 0,      # tổng items thiếu ngày
            })
            who = it["who"]
            if who and who not in entry["who_set"]:
                entry["who_set"].append(who)
            tname = (it["team_name"] or it["team_id"] or "").replace("team-", "Team ").title()
            if tname and tname not in entry["team_set"]:
                entry["team_set"].append(tname)
            entry["total"] += it["amount_vnd"]
            if it["for_date"]:
                k = it["for_date"].isoformat()
                entry["by_day"][k] = entry["by_day"].get(k, 0) + it["amount_vnd"]
            else:
                entry["missing"] += it["amount_vnd"]

        # Row 1: Start date | from | <span> | NGÂN SÁCH THÁNG | <total>
        grand_total = sum(e["total"] for e in agg.values())
        ws["A1"] = "Start date"; ws["A1"].fill = info_fill; ws["A1"].font = Font(bold=True); ws["A1"].alignment = center
        ws["B1"] = date_from.strftime("%d/%m/%Y"); ws["B1"].alignment = center
        ws["D1"] = "NGÂN SÁCH THÁNG"; ws["D1"].fill = info_fill; ws["D1"].font = Font(bold=True); ws["D1"].alignment = center
        ws["E1"] = grand_total; ws["E1"].fill = total_fill; ws["E1"].font = Font(bold=True, color="C00000", size=12)
        ws["E1"].number_format = '#,##0'; ws["E1"].alignment = right_align

        ws["A2"] = "End date"; ws["A2"].fill = info_fill; ws["A2"].font = Font(bold=True); ws["A2"].alignment = center
        ws["B2"] = date_to.strftime("%d/%m/%Y"); ws["B2"].alignment = center

        # Row 4: Header — STT | TK | THẺ | NHÂN VIÊN | TEAM | TỔNG | dates
        has_missing = any(e["missing"] > 0 for e in agg.values())
        headers = ["STT", "TÊN TÀI KHOẢN", "THẺ", "NHÂN VIÊN", "TEAM", "TỔNG CHẠY"] + [d.strftime("%d/%m/%Y") for d in days]
        if has_missing:
            headers.append("⚠ Thiếu ngày")
        for ci, h in enumerate(headers, start=1):
            c = ws.cell(row=4, column=ci, value=h)
            c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = border

        # Sort agg: TK có chứa 'TH+số' (cty) trước, rồi HKĐ; trong mỗi nhóm sort theo tên
        import re as _re
        _re_cty = _re.compile(r"th[0-9]", _re.IGNORECASE)
        def _sort_key(item):
            tk = item[1]["tk"]
            is_cty = bool(_re_cty.search(tk or ""))
            return (0 if is_cty else 1, tk.lower())

        # Data rows
        r = 5
        first_day_col = 7  # cột G (sau STT|TK|THẺ|NV|TEAM|TỔNG)
        for stt, (key, e) in enumerate(sorted(agg.items(), key=_sort_key), start=1):
            ws.cell(row=r, column=1, value=stt).alignment = center
            ws.cell(row=r, column=2, value=e["tk"]).font = Font(bold=bool(_re_cty.search(e["tk"] or "")))
            ws.cell(row=r, column=3, value=e["card"] or "").alignment = center
            ws.cell(row=r, column=4, value=", ".join(e["who_set"]) or "—")
            ws.cell(row=r, column=5, value=", ".join(e["team_set"]) or "—").alignment = center
            c_total = ws.cell(row=r, column=6, value=e["total"])
            c_total.number_format = '#,##0'; c_total.fill = money_fill; c_total.font = Font(bold=True)
            c_total.alignment = right_align
            for di, d_ in enumerate(days):
                amt = e["by_day"].get(d_.isoformat(), 0)
                cc = ws.cell(row=r, column=first_day_col+di, value=amt if amt else None)
                cc.number_format = '#,##0'; cc.alignment = right_align
            if has_missing:
                cc = ws.cell(row=r, column=first_day_col+len(days), value=e["missing"] if e["missing"] else None)
                cc.number_format = '#,##0'; cc.alignment = right_align
                if e["missing"]:
                    cc.fill = PatternFill("solid", fgColor="FCE4D6")
            for ci in range(1, len(headers)+1):
                ws.cell(row=r, column=ci).border = border
            r += 1

        # Column widths
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 22   # NV
        ws.column_dimensions["E"].width = 14   # TEAM
        ws.column_dimensions["F"].width = 14   # TỔNG
        for di in range(len(days)):
            ws.column_dimensions[get_column_letter(first_day_col+di)].width = 13
        if has_missing:
            ws.column_dimensions[get_column_letter(first_day_col+len(days))].width = 13

        # Freeze: giữ 6 cột đầu + 4 hàng header khi cuộn
        ws.freeze_panes = "G5"

    # Sort sheets theo tên
    wb._sheets.sort(key=lambda s: s.title)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname_team = (team_filter.replace("team-", "") + "_") if team_filter else "all_"
    fname = f"budget_{fname_team}{date_from.strftime('%Y%m%d')}_{date_to.strftime('%Y%m%d')}.xlsx"
    from flask import send_file
    return send_file(
        buf, as_attachment=True, download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _normalize_vn(s: str) -> str:
    """Chuẩn hoá tên: lowercase + bỏ dấu + bỏ ký tự đặc biệt."""
    import unicodedata as _ud
    s = _ud.normalize("NFD", s or "")
    s = "".join(c for c in s if _ud.category(c) != "Mn")
    s = s.lower()
    # Đổi ký tự đặc biệt thành space
    import re as _re
    s = _re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def _suggest_team_for_group(group_name: str, teams_all: list) -> str:
    """Fuzzy match tên Zalo group → team_code. Trả '' nếu không chắc."""
    if not group_name:
        return ""
    norm = _normalize_vn(group_name)
    # Bỏ các từ phổ biến không phải tên leader
    stopwords = {"kd", "team", "kinh", "doanh", "sai", "gon", "da", "nang",
                 "binh", "duong", "cn", "hcm", "vu", "1", "2", "3", "chi", "nhanh"}
    tokens = [t for t in norm.split() if t not in stopwords]
    if not tokens:
        return ""
    # So sánh với suffix team_code
    best, best_score = "", 0
    for t in teams_all:
        team_code = t.get("team_id", "")  # vd 'team-nam'
        team_suffix = _normalize_vn(team_code.replace("team-", ""))
        if not team_suffix:
            continue
        # Match: team_suffix là token hoặc xuất hiện trong tokens nối liền
        joined = "".join(tokens)
        score = 0
        if team_suffix in tokens:
            score = 100 + len(team_suffix)  # exact token
        elif team_suffix in joined:
            score = 50 + len(team_suffix)   # substring
        if score > best_score:
            best, best_score = team_code, score
    return best


@budget_chat_bp.route("/zalo-mapping")
@_login_required
def zalo_mapping():
    """ĐÃ GỘP sang trang "Lan gửi báo cáo cho ai?" (/chi-phi-qc/lan-bao-cao) — sếp 12/08.

    Trang cũ phục vụ chat ngân sách (module chưa dùng: 0 tin nhắn, 0 nhóm map) và
    trùng chức năng map Zalo↔NV. Giữ route để link/bookmark cũ không 404.
    """
    return redirect("/chi-phi-qc/lan-bao-cao")


def _zalo_mapping_cu():
    """Bản cũ — giữ nguyên code phòng khi cần bật lại."""
    user = _current_user()
    role = (user.get("role") or "").lower()
    if role not in _MANAGER_ROLES:
        abort(403)

    from db import get_conn
    from app_ctx import load_config
    with get_conn() as conn:
        with conn.cursor() as cur:
            # Pending senders chưa map
            cur.execute("""
                SELECT sender_uid, sender_name, last_thread_id, message_count,
                       first_seen, last_seen, last_body_snippet
                  FROM zalo_pending_senders
                 WHERE sender_uid NOT IN (SELECT zalo_uid FROM users WHERE zalo_uid IS NOT NULL)
                 ORDER BY last_seen DESC
                 LIMIT 200
            """)
            pending = [{
                "sender_uid": r[0], "sender_name": r[1], "last_thread_id": r[2],
                "count": int(r[3] or 0), "first_seen": r[4], "last_seen": r[5],
                "snippet": r[6] or "",
            } for r in cur.fetchall()]

            # Users đã map
            cur.execute("""
                SELECT id, username, full_name, zalo_uid,
                       (SELECT team_code FROM teams WHERE id = users.team_id) AS team_code
                  FROM users
                 WHERE zalo_uid IS NOT NULL
                 ORDER BY full_name
            """)
            mapped = [{
                "id": int(r[0]), "username": r[1], "full_name": r[2],
                "zalo_uid": r[3], "team_code": r[4],
            } for r in cur.fetchall()]

            # All users để dropdown chọn
            cur.execute("""
                SELECT u.id, u.username, u.full_name, t.team_code, t.team_name
                  FROM users u
                  LEFT JOIN teams t ON t.id = u.team_id
                 WHERE u.status='active'
                 ORDER BY t.team_name NULLS LAST, u.full_name
            """)
            all_users = [{
                "id": int(r[0]), "username": r[1], "full_name": r[2],
                "team_code": r[3], "team_name": r[4],
            } for r in cur.fetchall()]

    # Thread mapping từ app_config
    cfg = load_config()
    threads = []
    for k, v in cfg.items():
        if isinstance(k, str) and k.startswith("zalo_thread_"):
            threads.append({"thread_id": k[len("zalo_thread_"):], "team_code": str(v)})
    threads.sort(key=lambda x: x["team_code"])

    teams_all = _list_all_teams()

    # Auto-scan groups Lan đang trong (fail silent nếu bridge offline)
    auto_groups = []
    all_groups = []
    try:
        import os as _os, requests as _rq
        secret = (_os.environ.get("ZALO_BRIDGE_SECRET") or "").strip()
        if secret:
            url = _os.environ.get("ZALO_BRIDGE_OUTBOUND_URL", "http://127.0.0.1:5051/send").replace("/send", "/list-groups")
            r = _rq.get(url, headers={"X-Bridge-Secret": secret}, timeout=8)
            if r.status_code == 200:
                gs = r.json().get("groups", [])
                all_groups = [{"thread_id": g["thread_id"], "name": g.get("name", ""),
                               "member_count": g.get("member_count", 0)} for g in gs]
                # Loại bỏ groups đã map dưới mọi dạng (team budget, expense inbox CTY,
                # expense inbox RIÊNG hardcode). Không hiển thị "chưa map" giả.
                mapped_tids = {t["thread_id"] for t in threads}
                for k in cfg.keys():
                    if isinstance(k, str) and k.startswith("zalo_expense_thread_"):
                        mapped_tids.add(k[len("zalo_expense_thread_"):])
                # Hardcode private CP inbox (xem expense_chat._get_private_inbox_thread_id)
                try:
                    from modules.expense_chat import _get_private_inbox_thread_id
                    priv_tid = _get_private_inbox_thread_id()
                    if priv_tid:
                        mapped_tids.add(str(priv_tid))
                except Exception:
                    pass
                for g in gs:
                    if g["thread_id"] in mapped_tids:
                        continue
                    auto_groups.append({
                        "thread_id": g["thread_id"],
                        "name": g.get("name", ""),
                        "member_count": g.get("member_count", 0),
                        "suggested_team": _suggest_team_for_group(g.get("name", ""), teams_all),
                    })
    except Exception as exc:
        logger.info("zalo_mapping auto-scan skip: %s", exc)

    ads_report_tids = {t for t in str(cfg.get("lan_ads_report_threads") or "").split(",") if t.strip()}
    # nhóm đã bật nhưng bridge không thấy (VD bridge offline) → vẫn hiện để tắt được
    _seen = {g["thread_id"] for g in all_groups}
    for _tid in ads_report_tids:
        if _tid not in _seen:
            all_groups.append({"thread_id": _tid, "name": "(nhóm đã bật — bridge chưa quét thấy)", "member_count": 0})
    return render_template(
        "budget_chat/zalo_mapping.html",
        pending=pending, mapped=mapped, all_users=all_users,
        threads=threads, teams_all=teams_all,
        auto_groups=auto_groups, all_groups=all_groups,
        ads_report_tids=ads_report_tids,
        current_user=user,
    )


@budget_chat_bp.route("/zalo-mapping/assign-user", methods=["POST"])
@_login_required
def zalo_mapping_assign_user():
    user = _current_user()
    role = (user.get("role") or "").lower()
    if role not in _MANAGER_ROLES:
        abort(403)
    sender_uid = (request.form.get("sender_uid") or "").strip()
    target_user_id = request.form.get("user_id", "").strip()
    if not sender_uid:
        flash("Thiếu sender_uid", "warning")
        return redirect(url_for("budget_chat.zalo_mapping"))
    from db import get_conn
    with get_conn() as conn:
        with conn.cursor() as cur:
            if target_user_id and target_user_id != "0":
                cur.execute("UPDATE users SET zalo_uid = NULL WHERE zalo_uid = %s AND id <> %s",
                            (sender_uid, int(target_user_id)))
                cur.execute("UPDATE users SET zalo_uid = %s WHERE id = %s",
                            (sender_uid, int(target_user_id)))
                cur.execute("DELETE FROM zalo_pending_senders WHERE sender_uid = %s", (sender_uid,))
                flash(f"Đã map Zalo {sender_uid[:12]}… → user {target_user_id}", "success")
            else:
                # user_id=0 → unmap (xoá pending)
                cur.execute("DELETE FROM zalo_pending_senders WHERE sender_uid = %s", (sender_uid,))
                flash("Đã bỏ qua sender này", "info")
            conn.commit()
    return redirect(url_for("budget_chat.zalo_mapping"))


@budget_chat_bp.route("/zalo-mapping/unmap-user", methods=["POST"])
@_login_required
def zalo_mapping_unmap_user():
    user = _current_user()
    role = (user.get("role") or "").lower()
    if role not in _MANAGER_ROLES:
        abort(403)
    uid = int(request.form.get("user_id") or 0)
    if uid:
        from db import get_conn
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute("UPDATE users SET zalo_uid = NULL WHERE id = %s", (uid,))
                conn.commit()
        flash("Đã unmap", "info")
    return redirect(url_for("budget_chat.zalo_mapping"))


@budget_chat_bp.route("/zalo-mapping/toggle-ads-report", methods=["POST"])
@_login_required
def zalo_mapping_toggle_ads_report():
    """Bật/tắt 1 nhóm Zalo nhận BÁO CÁO ADS (Lãi/Lỗ) — lưu app_config lan_ads_report_threads."""
    user = _current_user()
    role = (user.get("role") or "").lower()
    if role not in _MANAGER_ROLES:
        abort(403)
    tid = (request.form.get("thread_id") or "").strip()
    enable = (request.form.get("enable") or "") == "1"
    if not tid:
        flash("Thiếu thread_id", "warning")
        return redirect(url_for("budget_chat.zalo_mapping"))
    from app_ctx import load_config, save_config_key
    cur_set = {t for t in str((load_config() or {}).get("lan_ads_report_threads") or "").split(",") if t.strip()}
    if enable:
        cur_set.add(tid)
        flash("Đã BẬT báo cáo Ads cho nhóm này ✅", "success")
    else:
        cur_set.discard(tid)
        flash("Đã tắt báo cáo Ads cho nhóm này", "info")
    save_config_key("lan_ads_report_threads", ",".join(sorted(cur_set)))
    return redirect(url_for("budget_chat.zalo_mapping"))


@budget_chat_bp.route("/zalo-mapping/assign-thread", methods=["POST"])
@_login_required
def zalo_mapping_assign_thread():
    user = _current_user()
    role = (user.get("role") or "").lower()
    if role not in _MANAGER_ROLES:
        abort(403)
    tid = (request.form.get("thread_id") or "").strip()
    team_code = (request.form.get("team_code") or "").strip()
    if not tid:
        flash("Thiếu thread_id", "warning")
        return redirect(url_for("budget_chat.zalo_mapping"))
    from app_ctx import save_config_key
    key = f"zalo_thread_{tid}"
    if team_code:
        save_config_key(key, team_code)
        flash(f"Đã map thread {tid[:14]}… → {team_code}", "success")
    else:
        # Empty team_code = remove mapping
        try:
            from db import get_conn
            with get_conn() as conn:
                with conn.cursor() as cur:
                    cur.execute("DELETE FROM app_config WHERE key = %s", (key,))
                    conn.commit()
            flash("Đã gỡ map thread", "info")
        except Exception as exc:
            flash(f"Lỗi gỡ map: {exc}", "danger")
    return redirect(url_for("budget_chat.zalo_mapping"))


@budget_chat_bp.route("/telegram-toggle", methods=["POST"])
@_login_required
def telegram_toggle():
    """BẬT/TẮT forward chat NS sang Telegram. Manager-only."""
    user = _current_user()
    role = (user.get("role") or "").lower()
    if role not in _MANAGER_ROLES:
        return jsonify({"error": "forbidden"}), 403
    enabled = (request.form.get("enabled") or "").strip() in ("1", "true", "on", "yes")
    try:
        from app_ctx import save_config_key
        save_config_key("budget_chat_telegram_enabled", "true" if enabled else "false")
        logger.info("telegram_toggle: user=%s set enabled=%s", user.get("username"), enabled)
        return jsonify({"ok": True, "enabled": enabled})
    except Exception as exc:
        logger.error("telegram_toggle error: %s", exc)
        return jsonify({"error": str(exc)}), 500


@budget_chat_bp.route("/items/<int:item_id>/edit", methods=["POST"])
@_login_required
def edit_item_route(item_id: int):
    user = _current_user()
    from db import get_conn
    from modules.budget_chat.repository import edit_item

    # Check ownership: NV chỉ sửa item của mình; manager/admin sửa tất cả
    role = (user.get("role") or "").lower()
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT user_id, team_id FROM budget_items WHERE id=%s AND deleted_at IS NULL", (item_id,))
            row = cur.fetchone()
            if not row:
                return jsonify({"error": "not found"}), 404
            owner_id, team_id = int(row[0]), row[1]
            if role not in _MANAGER_ROLES and int(user["id"]) != owner_id:
                return jsonify({"error": "forbidden"}), 403

            ok = edit_item(
                cur, item_id, int(user["id"]),
                tk_name_raw=request.form.get("tk_name_raw") or None,
                card_last4=request.form.get("card_last4"),
                amount_vnd=int(request.form.get("amount_vnd")) if request.form.get("amount_vnd") else None,
                for_date=request.form.get("for_date") or None,
            )
            conn.commit()
    return jsonify({"ok": ok})


@budget_chat_bp.route("/items/<int:item_id>/delete", methods=["POST"])
@_login_required
def delete_item_route(item_id: int):
    user = _current_user()
    from db import get_conn
    from modules.budget_chat.repository import soft_delete_item

    role = (user.get("role") or "").lower()
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT user_id FROM budget_items WHERE id=%s AND deleted_at IS NULL", (item_id,))
            row = cur.fetchone()
            if not row:
                return jsonify({"error": "not found"}), 404
            owner_id = int(row[0])
            if role not in _MANAGER_ROLES and int(user["id"]) != owner_id:
                return jsonify({"error": "forbidden"}), 403
            ok = soft_delete_item(cur, item_id, int(user["id"]))
            conn.commit()
    return jsonify({"ok": ok})


@budget_chat_bp.route("/messages/<int:msg_id>/lan-ack", methods=["POST"])
@_login_required
def lan_ack_after_edit(msg_id: int):
    """Sau khi NV save modal sửa items → Lan post 1 câu xác nhận đã ghi nhận."""
    user = _current_user()
    from db import get_conn
    from modules.budget_chat.repository import list_items_for_message, insert_lan_message
    from modules.budget_chat.lan_personality import ai_reply_ack_edit

    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT team_id, user_id FROM budget_chat_messages WHERE id=%s", (msg_id,))
            row = cur.fetchone()
            if not row:
                return jsonify({"error": "not found"}), 404
            team_id, msg_owner_id = row[0], int(row[1])
            # NV chỉ ack được tin của mình; manager ack được tất cả.
            role = (user.get("role") or "").lower()
            if role not in _MANAGER_ROLES and int(user["id"]) != msg_owner_id:
                return jsonify({"error": "forbidden"}), 403

            items = list_items_for_message(cur, msg_id)
            if not items:
                return jsonify({"ok": False, "reason": "no items"})

            # Lấy ngày từ item đầu (modal áp 1 ngày chung)
            for_date = items[0].get("for_date")
            for_date_vn = "(chưa rõ ngày)"
            if for_date:
                try:
                    from datetime import date as _date
                    d = _date.fromisoformat(for_date)
                    for_date_vn = d.strftime("%d/%m/%Y")
                except Exception:
                    for_date_vn = str(for_date)

            summary = ", ".join(
                f"TK {it.get('tk_name_raw','?')}"
                + (f" thẻ {it.get('card_last4')}" if it.get('card_last4') else "")
                + f" {int(it.get('amount_vnd') or 0):,}đ".replace(",", ".")
                for it in items[:5]
            )

            nv_name = user.get("full_name") or user.get("username") or "bạn"
            lan_text = ai_reply_ack_edit(nv_name, summary, for_date_vn)
            lan_id = insert_lan_message(cur, team_id, lan_text)
            conn.commit()
    # Tag chủ tin (msg_owner_id), không phải người sửa
    _push_lan_to_zalo(team_id, lan_text, user_id=msg_owner_id)
    return jsonify({"ok": True, "lan_message_id": lan_id})


@budget_chat_bp.route("/api/match-tk")
@_login_required
def api_match_tk():
    """Search TK QC để gán vào item khi NV sửa."""
    q = (request.args.get("q") or "").strip()
    if not q:
        return jsonify({"results": []})
    from modules.budget_chat.tk_matcher import _load_all_accounts
    accs = _load_all_accounts()
    ql = q.lower().replace(" ", "")
    matches = []
    for a in accs:
        if ql in a["short_norm"] or ql in a["account_name"].lower().replace(" ", ""):
            matches.append({
                "fb_ad_account_id": a["ad_account_id"],
                "short": a["short"],
                "name": a["account_name"],
            })
        if len(matches) >= 20:
            break
    return jsonify({"results": matches})


def register_budget_chat_module(app):
    """Đăng ký blueprint vào Flask app — gọi từ web_app.py."""
    app.register_blueprint(budget_chat_bp)
    logger.info("budget_chat module registered")
